#!/usr/bin/env python3
"""Export the final strict C4 candidate table to a compact Excel workbook."""

from __future__ import annotations

import argparse
import csv
import pathlib

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--table", type=pathlib.Path, required=True)
    parser.add_argument("--out", type=pathlib.Path, required=True)
    args = parser.parse_args()
    rows = list(csv.DictReader(args.table.open(), delimiter="\t"))
    fields = [
        "rank", "candidate", "n_mutations", "mutations", "mean_iptm", "median_iptm",
        "min_iptm", "mean_weakest_chain", "n_models", "full_connected_models",
        "unclashed_models", "no_interchain_ss_models", "all_native_c74_c121_models",
        "sequence",
    ]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "AF3严格四聚体20候选"
    sheet.append(fields)
    for row in rows:
        sheet.append([row.get(field, "") for field in fields])
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = {
        "A": 8, "B": 28, "C": 12, "D": 75, "E": 14, "F": 14, "G": 14,
        "H": 20, "I": 10, "J": 22, "K": 18, "L": 24, "M": 28, "N": 90,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in {4, 14})
    notes = workbook.create_sheet("判定说明")
    notes.append(["字段", "标准"])
    notes.append(["AF3", "5个独立seed × 每seed 5 samples，共25模型；算术平均ipTM严格大于0.80"])
    notes.append(["结构", "25/25四链连通且无clash"])
    notes.append(["化学", "禁止新增Cys；25/25无亚基间二硫键；保留SIINFEKL"])
    notes.append(["突变", "相对P3-13R严格少于25个"])
    notes.column_dimensions["A"].width = 16
    notes.column_dimensions["B"].width = 90
    args.out.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.out)
    print(f"wrote {len(rows)} candidates to {args.out}")


if __name__ == "__main__":
    main()
